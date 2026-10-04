from pathlib import Path

import pytest

import app.models  # noqa: F401  registers relationship string refs before we touch the mapper
from app.core.config import settings
from app.crud.attachment_text import (
    AttachmentExtractionFailedError,
    AttachmentNotExtractableError,
    convert_to_markdown,
    get_extracted_text,
    is_extractable,
    is_near_empty,
)
from app.models.attachment import Attachment
from app.models.enums import AttachmentKind

LECTURE_TEXT = (
    "Dijkstra's algorithm finds shortest paths from a single source in a graph "
    "with non-negative edge weights using a priority queue."
)


def write_pdf(path: Path, lines: list[str]) -> Path:
    """Write a minimal single-page PDF showing `lines` in Helvetica (no PDF library needed)."""
    text_ops = "".join(f"({line}) Tj 0 -16 Td " for line in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {text_ops}ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref_at)
    path.write_bytes(bytes(out))
    return path


def write_pptx(path: Path, title: str, body: str) -> Path:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    deck.save(str(path))
    return path


def make_attachment(kind: AttachmentKind, filename: str, file_path: str = "unused") -> Attachment:
    return Attachment(kind=kind, filename=filename, file_path=file_path)


# --- pure functions ---------------------------------------------------------


@pytest.mark.parametrize(
    ("kind", "filename", "expected"),
    [
        (AttachmentKind.pdf, "notes.pdf", True),
        (AttachmentKind.pptx, "slides.PPTX", True),
        (AttachmentKind.pptx, "legacy.ppt", False),
        (AttachmentKind.video, "lecture.mp4", False),
        (AttachmentKind.audio, "lecture.mp3", False),
        (AttachmentKind.image, "diagram.png", False),
        (AttachmentKind.other, "data.csv", False),
    ],
)
def test_is_extractable_filters_by_kind(kind, filename, expected):
    assert is_extractable(kind, filename) is expected


def test_is_near_empty_flags_blank_and_whitespace_only_text():
    assert is_near_empty("")
    assert is_near_empty("\n\n   \f\n")
    assert is_near_empty("# 1\n\n- -- 2 --")


def test_is_near_empty_accepts_real_text():
    assert not is_near_empty(LECTURE_TEXT)


# --- get_extracted_text -----------------------------------------------------


def test_get_extracted_text_converts_once_then_uses_cache():
    calls: list[str] = []

    def fake_convert(path: str) -> str:
        calls.append(path)
        return LECTURE_TEXT

    attachment = make_attachment(AttachmentKind.pdf, "notes.pdf", "submodules/1/x.pdf")

    first = get_extracted_text(attachment, convert=fake_convert)
    second = get_extracted_text(attachment, convert=fake_convert)

    assert calls == [str(Path(settings.upload_dir) / "submodules" / "1" / "x.pdf")]  # read from the upload folder
    assert attachment.extracted_markdown == LECTURE_TEXT
    assert first == second
    assert first.markdown == LECTURE_TEXT
    assert not first.near_empty


def test_get_extracted_text_caches_an_empty_result():
    calls: list[str] = []

    def fake_convert(path: str) -> str:
        calls.append(path)
        return ""

    attachment = make_attachment(AttachmentKind.pdf, "scan.pdf")

    get_extracted_text(attachment, convert=fake_convert)
    result = get_extracted_text(attachment, convert=fake_convert)

    assert len(calls) == 1
    assert attachment.extracted_markdown == ""
    assert result.near_empty


def test_get_extracted_text_only_flags_near_empty_for_pdfs():
    attachment = make_attachment(AttachmentKind.pptx, "slides.pptx")

    result = get_extracted_text(attachment, convert=lambda _: "Title")

    assert not result.near_empty


@pytest.mark.parametrize(
    ("kind", "filename"),
    [(AttachmentKind.video, "lecture.mp4"), (AttachmentKind.image, "x.png"), (AttachmentKind.pptx, "old.ppt")],
)
def test_get_extracted_text_rejects_unsupported_kinds_without_converting(kind, filename):
    def fail(_: str) -> str:
        raise AssertionError("conversion should not run")

    attachment = make_attachment(kind, filename)

    with pytest.raises(AttachmentNotExtractableError, match="not supported"):
        get_extracted_text(attachment, convert=fail)
    assert attachment.extracted_markdown is None


# --- real conversion with markitdown -----------------------------------------


def test_converts_pdf_with_text(tmp_path):
    pdf = write_pdf(tmp_path / "lecture.pdf", ["Shortest paths", LECTURE_TEXT])
    attachment = make_attachment(AttachmentKind.pdf, "lecture.pdf", str(pdf))

    result = get_extracted_text(attachment)

    assert "Shortest paths" in result.markdown
    assert "priority queue" in result.markdown
    assert not result.near_empty


def test_flags_near_empty_pdf(tmp_path):
    pdf = write_pdf(tmp_path / "scan.pdf", ["3"])
    attachment = make_attachment(AttachmentKind.pdf, "scan.pdf", str(pdf))

    result = get_extracted_text(attachment)

    assert result.near_empty
    assert attachment.extracted_markdown is not None


def test_converts_pptx(tmp_path):
    pptx = write_pptx(tmp_path / "slides.pptx", "Graph Traversal", "Breadth-first search uses a queue")
    attachment = make_attachment(AttachmentKind.pptx, "slides.pptx", str(pptx))

    result = get_extracted_text(attachment)

    assert "Graph Traversal" in result.markdown
    assert "Breadth-first search uses a queue" in result.markdown
    assert not result.near_empty


def test_failed_conversion_raises_and_is_not_cached(tmp_path):
    attachment = make_attachment(AttachmentKind.pdf, "gone.pdf", str(tmp_path / "missing.pdf"))

    with pytest.raises(AttachmentExtractionFailedError):
        get_extracted_text(attachment, convert=convert_to_markdown)
    assert attachment.extracted_markdown is None
